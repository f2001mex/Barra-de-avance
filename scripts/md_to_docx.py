from pathlib import Path
from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

ROOT = Path(__file__).resolve().parents[1]
DOCS = ROOT / 'docs'


def set_style(doc):
    style = doc.styles['Normal']
    style.font.name = 'Calibri'
    style._element.rPr.rFonts.set(qn('w:eastAsia'), 'Calibri')
    style.font.size = Pt(10)


def add_footer(section):
    p = section.footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fld = OxmlElement('w:fldSimple')
    fld.set(qn('w:instr'), 'PAGE')
    p.add_run()._r.append(fld)


def add_code(doc, text):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.font.name = 'Consolas'
    r._element.rPr.rFonts.set(qn('w:eastAsia'), 'Consolas')
    r.font.size = Pt(9)


def convert(md_path: Path):
    lines = md_path.read_text(encoding='utf-8').splitlines()
    doc = Document()
    set_style(doc)
    sec = doc.sections[0]
    sec.top_margin = Inches(0.7)
    sec.bottom_margin = Inches(0.7)
    sec.left_margin = Inches(0.8)
    sec.right_margin = Inches(0.8)
    add_footer(sec)
    in_code = False
    code_lines = []
    for line in lines:
        if line.strip().startswith('```'):
            if in_code:
                add_code(doc, '\n'.join(code_lines))
                code_lines = []
                in_code = False
            else:
                in_code = True
            continue
        if in_code:
            code_lines.append(line)
            continue
        if not line.strip():
            doc.add_paragraph('')
            continue
        if line.startswith('# '):
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = p.add_run(line[2:])
            r.bold = True
            r.font.size = Pt(18)
            continue
        if line.startswith('## '):
            doc.add_heading(line[3:], level=1)
            continue
        if line.startswith('### '):
            doc.add_heading(line[4:], level=2)
            continue
        if line.startswith('- '):
            doc.add_paragraph(line[2:], style='List Bullet')
            continue
        if line[:3].isdigit() and line[3:5] == '. ':
            doc.add_paragraph(line[5:], style='List Number')
            continue
        if len(line) > 2 and line[0].isdigit() and line[1:3] == '. ':
            doc.add_paragraph(line[3:], style='List Number')
            continue
        if line.startswith('|') and line.endswith('|'):
            continue
        doc.add_paragraph(line)
    out = md_path.with_suffix('.docx')
    doc.save(out)
    print(out)


for md in sorted(DOCS.glob('*2026-03-13.md')):
    convert(md)
